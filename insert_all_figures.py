import docx
import shutil
import os
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

orig_doc_path = 'c:/Users/Yurec/OneDrive/НМК по дисциплінам/КСМ/НМК з КСМ/Лабораторні роботи дистанційно/ЛАБОРАТОРНА РОБОТА № 1 Основи побудови комп’ютерних мереж.docx'
backup_path = orig_doc_path + '.bak_before_figs'
screenshots_dir = 'c:/Users/Yurec/OneDrive/НМК по дисциплінам/КСМ/НМК з КСМ/Лабораторні роботи дистанційно/rj45_simulator/screenshots'

# 1. Backup original document
shutil.copy2(orig_doc_path, backup_path)
print(f"Backup created: {backup_path}")

doc = docx.Document(orig_doc_path)

# Define figures to insert
figures_to_insert = [
    {
        'before_prefix': '1.1. Дослідження та віртуальний',
        'img': os.path.join(screenshots_dir, 'fig_1_9_theory_guide.png'),
        'caption': 'Рис. 1.9. Вбудований інтерактивний довідник стандартів монтажу TIA-568 із колірними схемами та призначенням пар',
        'width': Inches(5.1)
    },
    {
        'before_prefix': '1.2. Дослідження та віртуальний',
        'img': os.path.join(screenshots_dir, 'fig_1_10_straight_t568b.png'),
        'caption': 'Рис. 1.10. Інтерфейс симулятора «RJ-45 Crimp Lab Pro»: успішний двосторонній монтаж та апаратне тестування прямого кабелю T568B у LAN-тестері (1000 Mbps Full Duplex)',
        'width': Inches(5.3)
    },
    {
        'before_prefix': '1.2. Дослідження та віртуальний',
        'img': os.path.join(screenshots_dir, 'fig_1_13_verification_act.png'),
        'caption': 'Рис. 1.11. Електронний «Акт верифікації та тестування лінку» із захисним цифровим хеш-кодом звіту',
        'width': Inches(4.8)
    },
    {
        'before_prefix': '1.4. Дослідження схем трасування',
        'img': os.path.join(screenshots_dir, 'fig_1_11_wiremap_crossover.png'),
        'caption': 'Рис. 1.12. Інтерактивна векторна схема трасування кабелю (Wiremap Crossover T568B ↔ T568A) зі схрещенням передавальних та прийомних ліній',
        'width': Inches(5.3)
    },
    {
        'before_prefix': '1.5. Порівняльний аналіз стандартів',
        'img': os.path.join(screenshots_dir, 'fig_1_12_defect_tester.png'),
        'caption': 'Рис. 1.13. Апаратна діагностика дефекту допоміжних пар у LAN-тестері: фіксація розриву гігабітного лінку та деградація швидкості до 100 Мбіт/с (Fast Ethernet)',
        'width': Inches(5.3)
    }
]

# We should insert in reverse or track paragraph references
for fig in figures_to_insert:
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith(fig['before_prefix']):
            target_idx = i
            break
    
    if target_idx is None:
        raise ValueError(f"Target prefix not found: {fig['before_prefix']}")
    
    p_target = doc.paragraphs[target_idx]
    
    # Insert image paragraph
    p_img = p_target.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.line_spacing = 1.0
    r_img = p_img.add_run()
    r_img.add_picture(fig['img'], width=fig['width'])
    
    # Insert caption paragraph
    p_cap = p_target.insert_paragraph_before()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    p_cap.paragraph_format.line_spacing = 1.15
    r_cap = p_cap.add_run(fig['caption'])
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    
    print(f"Inserted: {fig['caption'][:60]}... before paragraph '{fig['before_prefix']}'")

doc.save(orig_doc_path)
print(f"Document saved successfully! New size: {os.path.getsize(orig_doc_path)} bytes")
